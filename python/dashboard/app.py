import sys
from pathlib import Path

# Позволяет запускать:
# streamlit run python/dashboard/app.py
PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import pandas as pd
import streamlit as st
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

from sqlalchemy.exc import IntegrityError

from python.dashboard.data import (
    load_trains,
    load_passengers,
    load_luggage,
    get_train_options,
    get_passenger_options,
    create_train,
    update_train,
    delete_train,
    create_passenger,
    update_passenger,
    delete_passenger,
    create_luggage,
    update_luggage,
    delete_luggage,
    get_dashboard_summary,
    get_route_statistics,
    get_daily_statistics,
    get_carriage_statistics,
)

from python.dashboard.charts import (
    revenue_by_train,
    tickets_by_date,
    revenue_by_route,
)


st.set_page_config(
    page_title="Railway Analytics",
    page_icon="🚆",
    layout="wide",
    initial_sidebar_state="expanded",
)


# ============================================================
# CSS
# ============================================================

st.markdown(
    """
    <style>
        .main {
            background-color: #f5f7fa;
        }

        .block-container {
            padding-top: 2rem;
            padding-bottom: 2rem;
        }

        .dashboard-title {
            font-size: 36px;
            font-weight: 700;
            margin-bottom: 5px;
        }

        .dashboard-subtitle {
            color: #6b7280;
            font-size: 16px;
            margin-bottom: 25px;
        }

        div[data-testid="stMetric"] {
            background-color: white;
            border: 1px solid #e5e7eb;
            border-radius: 12px;
            padding: 18px;
        }

        .section-title {
            font-size: 22px;
            font-weight: 650;
            margin-top: 20px;
            margin-bottom: 10px;
        }

        /* Все формы ввода выполнены в единой тёмной цветовой гамме. */
        div[data-testid="stForm"] {
            background-color: #181a22;
            padding: 20px;
            border-radius: 12px;
            border: 1px solid #2d303a;
        }

        /* Поля внутри формы имеют тот же фон, без белых областей. */
        div[data-testid="stForm"] input,
        div[data-testid="stForm"] textarea,
        div[data-testid="stForm"] [data-baseweb="select"] > div,
        div[data-testid="stForm"] [data-baseweb="input"] > div {
            background-color: #272933 !important;
            color: #f3f4f6 !important;
            border-color: #3a3d48 !important;
        }

        div[data-testid="stForm"] label,
        div[data-testid="stForm"] label p {
            color: #f3f4f6 !important;
        }

        div[data-testid="stForm"] button {
            background-color: #171922;
            color: #f3f4f6;
            border-color: #3a3d48;
        }

        /* Кнопка сортировки находится на одной линии с панелью фильтрации. */
        div[data-testid="column"] button[kind="secondary"] {
            min-height: 38px;
        }
    </style>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# CACHE
# ============================================================

@st.cache_data(ttl=60)
def cached_trains():
    return load_trains()


@st.cache_data(ttl=60)
def cached_passengers():
    return load_passengers()


@st.cache_data(ttl=60)
def cached_luggage():
    return load_luggage()


@st.cache_data(ttl=60)
def cached_train_options():
    return get_train_options()


@st.cache_data(ttl=60)
def cached_passenger_options():
    return get_passenger_options()


@st.cache_data(ttl=60)
def cached_summary():
    return get_dashboard_summary()


@st.cache_data(ttl=60)
def cached_routes():
    return get_route_statistics()


@st.cache_data(ttl=60)
def cached_daily():
    return get_daily_statistics()


@st.cache_data(ttl=60)
def cached_carriages():
    return get_carriage_statistics()


def refresh_data():
    st.cache_data.clear()
    st.rerun()


# ============================================================
# EXCEL-ОТЧЁТЫ
# ============================================================

RUSSIAN_COLUMN_NAMES = {
    "id": "ID",
    "train_id": "ID поезда",
    "passenger_id": "ID пассажира",
    "train_number": "Номер поезда",
    "departure_date": "Дата отправления",
    "departure_station": "Станция отправления",
    "destination_station": "Станция назначения",
    "sold_tickets": "Продано билетов",
    "ticket_price": "Цена билета, BYN",
    "carriage_type": "Тип вагона",
    "total_revenue": "Выручка, BYN",
    "full_name": "ФИО пассажира",
    "luggage_count": "Количество багажа",
    "extra_places": "Дополнительные места",
    "train_count": "Количество поездов",
    "total_tickets": "Всего билетов",
    "average_ticket_price": "Средняя цена билета, BYN",
    "average_price": "Средняя цена, BYN",
    "route": "Маршрут",
    "date": "Дата",
    "count": "Количество",
    "revenue": "Выручка, BYN",
    "tickets": "Билеты",
    "total_luggage": "Всего багажа",
}

CURRENCY_COLUMNS = {
    "ticket_price",
    "total_revenue",
    "average_ticket_price",
    "average_price",
    "revenue",
}
DATE_COLUMNS = {"departure_date", "date"}


def prepare_excel_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    """Подготавливает данные для совместимого XLSX-отчёта."""
    if df is None:
        return pd.DataFrame()

    export_df = df.copy()

    # Перед экспортом восстанавливаем сведения о поезде по train_id.
    export_df = enrich_train_fields(export_df)

    # Приводим даты к настоящим датам Excel.
    for column in DATE_COLUMNS:
        if column in export_df.columns:
            export_df[column] = pd.to_datetime(
                export_df[column], errors="coerce"
            )

    # Все технические имена колонок переводим на русский.
    export_df = export_df.rename(
        columns={
            column: RUSSIAN_COLUMN_NAMES.get(column, column)
            for column in export_df.columns
        }
    )

    return export_df


def _excel_value(value):
    """Преобразует pandas/numpy-значение в безопасный тип для openpyxl."""
    if value is None:
        return None

    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass

    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()

    # numpy.int64/float64/bool и другие scalar-типы.
    if hasattr(value, "item"):
        try:
            value = value.item()
        except (ValueError, TypeError):
            pass

    # Остаточные даты/время оставляем как есть, строки и числа безопасны.
    return value


def _safe_sheet_name(name: str, used_names: set[str]) -> str:
    """Возвращает допустимое и уникальное имя листа Excel."""
    invalid = '<>:/\\|?*[]'
    cleaned = "".join("_" if char in invalid else char for char in str(name))
    cleaned = cleaned.strip() or "Отчёт"
    cleaned = cleaned[:31]

    candidate = cleaned
    number = 2
    while candidate in used_names:
        suffix = f"_{number}"
        candidate = f"{cleaned[:31-len(suffix)]}{suffix}"
        number += 1

    used_names.add(candidate)
    return candidate


def _write_excel_sheet(ws, df: pd.DataFrame, title: str) -> None:
    """Создаёт обычный, максимально совместимый с Excel отчёт."""
    export_df = prepare_excel_dataframe(df)
    columns = list(export_df.columns)
    column_count = max(len(columns), 1)

    # Даты пассажиров/багажа должны быть восстановлены до записи в Excel.
    # Ничего из исходных строк здесь не отбрасывается.

    # ------------------------------------------------------------
    # Заголовок отчёта
    # ------------------------------------------------------------
    ws.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=column_count,
    )
    title_cell = ws.cell(row=1, column=1, value=str(title))
    title_cell.font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill("solid", fgColor="172033")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    # Строка 2 специально оставляется пустой.
    ws.row_dimensions[2].height = 8

    # ------------------------------------------------------------
    # Заголовки таблицы
    # ------------------------------------------------------------
    header_row = 3
    header_fill = PatternFill("solid", fgColor="2F4057")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin = Side(style="thin", color="A6AEB8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_idx, column in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=str(column))
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = border

    ws.row_dimensions[header_row].height = 38

    # ------------------------------------------------------------
    # Данные
    # ------------------------------------------------------------
    original_columns = list(df.columns)
    first_data_row = header_row + 1

    for row_idx, row in enumerate(
        export_df.itertuples(index=False, name=None),
        start=first_data_row,
    ):
        for col_idx, raw_value in enumerate(row, start=1):
            value = _excel_value(raw_value)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = Font(name="Calibri", size=11)
            cell.border = border
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )

            original_column = original_columns[col_idx - 1]

            if original_column in DATE_COLUMNS and value is not None:
                cell.number_format = "dd.mm.yyyy"
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                )
            elif original_column in CURRENCY_COLUMNS and value is not None:
                cell.number_format = '#,##0.00 "BYN"'
                cell.alignment = Alignment(
                    horizontal="right",
                    vertical="center",
                )
            elif isinstance(value, (int, float)) and not isinstance(value, bool):
                cell.alignment = Alignment(
                    horizontal="right",
                    vertical="center",
                )

        # Чередование строк без использования сложных Excel-таблиц.
        if (row_idx - first_data_row) % 2 == 1:
            for col_idx in range(1, column_count + 1):
                ws.cell(row=row_idx, column=col_idx).fill = PatternFill(
                    "solid", fgColor="F3F6FA"
                )

        ws.row_dimensions[row_idx].height = 25

    last_row = max(header_row, ws.max_row)
    last_col = column_count

    # Если данных нет, всё равно оставляем красиво оформленную шапку.
    if len(export_df) == 0:
        ws.cell(row=header_row + 1, column=1, value="Нет данных для отображения")
        ws.cell(row=header_row + 1, column=1).font = Font(
            name="Calibri", size=11, italic=True, color="666666"
        )
        ws.cell(row=header_row + 1, column=1).border = border
        last_row = header_row + 1

    # ------------------------------------------------------------
    # Ширина столбцов
    # ------------------------------------------------------------
    for col_idx in range(1, last_col + 1):
        letter = get_column_letter(col_idx)
        max_length = len(str(ws.cell(header_row, col_idx).value or ""))

        for row_idx in range(first_data_row, last_row + 1):
            value = ws.cell(row_idx, col_idx).value
            if value is None:
                continue
            text_value = str(value)
            max_length = max(
                max_length,
                max((len(line) for line in text_value.splitlines()), default=0),
            )

        # Не даём длинному тексту растянуть весь отчёт.
        if max_length <= 12:
            width = 14
        elif max_length <= 20:
            width = max_length + 3
        elif max_length <= 35:
            width = max_length + 2
        else:
            width = 38

        ws.column_dimensions[letter].width = min(max(width, 12), 38)

    # ------------------------------------------------------------
    # Фильтр, закрепление, печать
    # ------------------------------------------------------------
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(last_col)}{last_row}"
    )
    ws.sheet_view.showGridLines = False

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5
    ws.print_title_rows = "1:3"
    ws.print_area = (
        f"A1:{get_column_letter(last_col)}{last_row}"
    )


def dataframe_to_excel(df: pd.DataFrame, sheet_name: str = "Отчёт") -> bytes:
    """Создаёт надёжный XLSX-файл с одним русскоязычным отчётом."""
    output = BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = _safe_sheet_name(sheet_name, set())
    _write_excel_sheet(ws, df, sheet_name)
    wb.save(output)
    return output.getvalue()


def dataframes_to_excel(sheets: dict[str, pd.DataFrame]) -> bytes:
    """Создаёт надёжный XLSX-файл с несколькими русскоязычными листами."""
    output = BytesIO()
    wb = Workbook()
    used_names: set[str] = set()

    if not sheets:
        ws = wb.active
        ws.title = "Отчёт"
        _write_excel_sheet(ws, pd.DataFrame(), "Отчёт")
    else:
        first = True
        for sheet_name, df in sheets.items():
            if first:
                ws = wb.active
                first = False
            else:
                ws = wb.create_sheet()

            safe_name = _safe_sheet_name(str(sheet_name), used_names)
            ws.title = safe_name
            _write_excel_sheet(ws, df, str(sheet_name))

    wb.save(output)
    return output.getvalue()


def excel_workbook_download(
    sheets: dict[str, pd.DataFrame],
    file_name: str,
    label: str,
) -> None:
    """Кнопка скачивания полноценного XLSX-отчёта."""
    st.download_button(
        label,
        data=dataframes_to_excel(sheets),
        file_name=file_name,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=False,
    )


def excel_download(
    df: pd.DataFrame,
    file_name: str,
    sheet_name: str,
    label: str = "Скачать Excel",
) -> None:
    """Кнопка скачивания полноценного XLSX-отчёта."""
    st.download_button(
        label,
        data=dataframe_to_excel(df, sheet_name),
        file_name=file_name,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=False,
    )


# ============================================================
# ФОРМАТИРОВАНИЕ
# ============================================================

def format_train_label(row) -> str:
    date_value = pd.to_datetime(row["departure_date"], errors="coerce")

    if pd.isna(date_value):
        date_text = "дата не указана"
    else:
        date_text = date_value.strftime("%d.%m.%Y")

    return (
        f'{row["train_number"]} — '
        f'{row["departure_station"]} → '
        f'{row["destination_station"]} — '
        f'{date_text}'
    )


def format_passenger_label(row) -> str:
    date_value = pd.to_datetime(row["departure_date"], errors="coerce")

    if pd.isna(date_value):
        date_text = "дата не указана"
    else:
        date_text = date_value.strftime("%d.%m.%Y")

    return (
        f'{row["full_name"]} — '
        f'{row["train_number"]} — '
        f'{row["departure_station"]} → '
        f'{row["destination_station"]} — '
        f'{date_text}'
    )


def display_date_columns(df: pd.DataFrame) -> pd.DataFrame:
    result = df.copy()

    if "departure_date" in result.columns:
        result["departure_date"] = pd.to_datetime(
            result["departure_date"],
            errors="coerce",
        ).dt.strftime("%d.%m.%Y")

    return result


def enrich_train_fields(df: pd.DataFrame) -> pd.DataFrame:
    """
    Гарантирует, что у пассажиров и багажа заполнены данные выбранного поезда.

    Источник истины для даты, номера поезда и маршрута — таблица trains.
    Если в исходном DataFrame какое-либо из этих значений отсутствует,
    оно восстанавливается по train_id.
    """
    if df is None or df.empty or "train_id" not in df.columns:
        return df.copy() if df is not None else pd.DataFrame()

    result = df.copy()

    try:
        train_df = load_trains().copy()
    except Exception:
        return result

    required = {
        "id",
        "train_number",
        "departure_date",
        "departure_station",
        "destination_station",
    }
    if not required.issubset(train_df.columns):
        return result

    train_df = train_df[
        [
            "id",
            "train_number",
            "departure_date",
            "departure_station",
            "destination_station",
        ]
    ].copy()

    train_df = train_df.rename(
        columns={
            "id": "_lookup_train_id",
            "train_number": "_lookup_train_number",
            "departure_date": "_lookup_departure_date",
            "departure_station": "_lookup_departure_station",
            "destination_station": "_lookup_destination_station",
        }
    )

    result["_merge_train_id"] = pd.to_numeric(
        result["train_id"], errors="coerce"
    )
    train_df["_merge_train_id"] = pd.to_numeric(
        train_df["_lookup_train_id"], errors="coerce"
    )

    result = result.merge(
        train_df,
        on="_merge_train_id",
        how="left",
        sort=False,
    )

    field_pairs = [
        ("train_number", "_lookup_train_number"),
        ("departure_date", "_lookup_departure_date"),
        ("departure_station", "_lookup_departure_station"),
        ("destination_station", "_lookup_destination_station"),
    ]

    for target, source in field_pairs:
        if source not in result.columns:
            continue

        if target not in result.columns:
            result[target] = result[source]
        else:
            # Заполняем только отсутствующие значения и не трогаем
            # уже корректно загруженные данные.
            result[target] = result[target].where(
                result[target].notna(),
                result[source],
            )

    helper_columns = [
        "_merge_train_id",
        "_lookup_train_id",
        "_lookup_train_number",
        "_lookup_departure_date",
        "_lookup_departure_station",
        "_lookup_destination_station",
    ]
    result = result.drop(
        columns=[c for c in helper_columns if c in result.columns]
    )

    # Возвращаем исходный порядок колонок.
    original_columns = list(df.columns)
    for column in [
        "train_number",
        "departure_date",
        "departure_station",
        "destination_station",
    ]:
        if column in result.columns and column not in original_columns:
            original_columns.append(column)

    return result[[c for c in original_columns if c in result.columns]]


def safe_contains(series: pd.Series, value: str) -> pd.Series:
    return (
        series.astype(str)
        .str.contains(value, case=False, na=False, regex=False)
    )


def handle_db_error(error: Exception, action: str):
    message = str(error)

    if "uq_train_number_date" in message:
        st.error(
            "Нельзя добавить или изменить поезд: "
            "такой номер поезда уже существует на указанную дату."
        )
    elif "fk_passenger_train" in message:
        st.error("Выбранный поезд не существует.")
    elif "fk_luggage_passenger" in message:
        st.error("Выбранный пассажир не существует.")
    else:
        st.error(f"Ошибка при операции «{action}»: {message}")


# ============================================================
# SIDEBAR
# ============================================================

with st.sidebar:
    st.title("🚆 Railway Analytics")
    st.caption("Система аналитики железнодорожных перевозок")
    st.divider()

    page = st.radio(
        "Навигация",
        [
            "Главная",
            "Поезда",
            "Пассажиры",
            "Багаж",
            "Маршруты",
            "Аналитика",
        ],
    )

    st.divider()

    if st.button(
        "⟳ Обновить данные",
        use_container_width=True,
    ):
        refresh_data()


trains = cached_trains()
passengers = enrich_train_fields(cached_passengers())
luggage = enrich_train_fields(cached_luggage())
summary = cached_summary()


# ============================================================
# HEADER
# ============================================================

st.markdown(
    '<div class="dashboard-title">Railway Analytics</div>',
    unsafe_allow_html=True,
)

st.markdown(
    '<div class="dashboard-subtitle">'
    'Аналитическая система железнодорожных перевозок'
    '</div>',
    unsafe_allow_html=True,
)


# ============================================================
# ГЛАВНАЯ
# ============================================================

if page == "Главная":

    st.markdown(
        '<div class="section-title">Основные показатели</div>',
        unsafe_allow_html=True,
    )

    if not trains.empty:
        col1, col2 = st.columns(2)

        with col1:
            st.pyplot(
                revenue_by_train(trains),
                use_container_width=True,
            )

        with col2:
            st.pyplot(
                tickets_by_date(trains),
                use_container_width=True,
            )

        st.markdown(
            '<div class="section-title">Выручка по маршрутам</div>',
            unsafe_allow_html=True,
        )

        st.pyplot(
            revenue_by_route(trains),
            use_container_width=True,
        )

    st.markdown(
        '<div class="section-title">Последние данные</div>',
        unsafe_allow_html=True,
    )

    st.dataframe(
        display_date_columns(trains.head(10)),
        use_container_width=True,
        hide_index=True,
    )


# ============================================================
# ПОЕЗДА
# ============================================================

elif page == "Поезда":

    st.header("Поезда")

    tab_view, tab_add, tab_edit, tab_delete = st.tabs(
        ["Просмотр", "Добавить", "Изменить", "Удалить"]
    )

    with tab_view:

        if trains.empty:
            st.info("В таблице поездов пока нет данных.")
        else:
            col1, col2 = st.columns([2, 1])

            with col1:
                search = st.text_input(
                    "Поиск",
                    placeholder=(
                        "Номер, станция отправления, "
                        "станция назначения, тип вагона..."
                    ),
                    key="train_search",
                )

            with col2:
                carriage_options = ["Все"] + sorted(
                    trains["carriage_type"]
                    .dropna()
                    .astype(str)
                    .unique()
                    .tolist()
                )

                carriage = st.selectbox(
                    "Тип вагона",
                    carriage_options,
                    key="train_carriage_filter",
                )

            valid_dates = trains["departure_date"].dropna()

            date_from = None
            date_to = None

            if not valid_dates.empty:
                min_date = valid_dates.min().date()
                max_date = valid_dates.max().date()

                col1, col2 = st.columns(2)

                with col1:
                    date_from = st.date_input(
                        "Дата от",
                        value=min_date,
                        min_value=min_date,
                        max_value=max_date,
                        key="train_date_from",
                    )

                with col2:
                    date_to = st.date_input(
                        "Дата до",
                        value=max_date,
                        min_value=min_date,
                        max_value=max_date,
                        key="train_date_to",
                    )

            filtered = trains.copy()

            if search.strip():
                mask = (
                    safe_contains(
                        filtered["train_number"],
                        search.strip(),
                    )
                    | safe_contains(
                        filtered["departure_station"],
                        search.strip(),
                    )
                    | safe_contains(
                        filtered["destination_station"],
                        search.strip(),
                    )
                    | safe_contains(
                        filtered["carriage_type"],
                        search.strip(),
                    )
                )

                filtered = filtered[mask]

            if carriage != "Все":
                filtered = filtered[
                    filtered["carriage_type"] == carriage
                ]

            if date_from is not None and date_to is not None:
                if date_from > date_to:
                    st.warning(
                        "Дата начала периода не может быть позже даты окончания."
                    )
                else:
                    date_from_ts = pd.Timestamp(date_from)
                    date_to_ts = pd.Timestamp(date_to)

                    filtered["departure_date"] = pd.to_datetime(
                        filtered["departure_date"],
                        errors="coerce",
                    )

                    filtered = filtered[
                        filtered["departure_date"].between(
                            date_from_ts,
                            date_to_ts,
                            inclusive="both",
                        )
                    ]

            # На странице «Поезда» записи всегда отображаются по ID.
            filtered = filtered.sort_values(
                by="id",
                ascending=True,
            )

            st.write(f"Найдено поездов: **{len(filtered)}**")

            display_df = display_date_columns(filtered)

            st.dataframe(
                display_df,
                use_container_width=True,
                hide_index=True,
            )

            excel_download(
                display_df,
                "trains_report.xlsx",
                "Поезда",
            )

    with tab_add:

        st.subheader("Добавление поезда")

        with st.form("add_train_form"):

            col1, col2 = st.columns(2)

            with col1:
                train_number = st.text_input(
                    "Номер поезда",
                    placeholder="Например, 900Б",
                )

                departure_date = st.date_input(
                    "Дата отправления",
                )

                departure_station = st.selectbox(
                    "Станция отправления",
                    [
                        "Минск",
                        "Брест",
                        "Витебск",
                        "Гомель",
                        "Москва",
                    ],
                )

            with col2:
                destination_station = st.selectbox(
                    "Станция назначения",
                    [
                        "Москва",
                        "Брест",
                        "Витебск",
                        "Гомель",
                        "Минск",
                    ],
                )

                sold_tickets = st.number_input(
                    "Продано билетов",
                    min_value=0,
                    value=0,
                    step=1,
                )

                ticket_price = st.number_input(
                    "Цена билета",
                    min_value=0.01,
                    value=30.00,
                    step=0.50,
                )

            carriage_type = st.selectbox(
                "Тип вагона",
                ["Плацкарт", "Купейный", "СВ", "Общий"],
            )

            submitted = st.form_submit_button(
                "Добавить поезд",
                use_container_width=True,
            )

            if submitted:

                if not train_number.strip():
                    st.warning("Введите номер поезда.")

                elif departure_station == destination_station:
                    st.warning(
                        "Станция отправления и станция назначения "
                        "должны различаться."
                    )

                else:
                    try:
                        create_train(
                            train_number=train_number,
                            departure_date=departure_date,
                            departure_station=departure_station,
                            destination_station=destination_station,
                            sold_tickets=sold_tickets,
                            ticket_price=ticket_price,
                            carriage_type=carriage_type,
                        )

                        st.success("Поезд успешно добавлен.")
                        refresh_data()

                    except IntegrityError as error:
                        handle_db_error(error, "добавление поезда")

    with tab_edit:

        if trains.empty:
            st.info("Нет поездов для изменения.")
        else:
            train_ids = trains["id"].astype(int).tolist()

            selected_id = st.selectbox(
                "Выберите поезд",
                train_ids,
                format_func=lambda x: format_train_label(
                    trains[trains["id"] == x].iloc[0]
                ),
                key="edit_train_id",
            )

            row = trains[trains["id"] == selected_id].iloc[0]

            with st.form("edit_train_form"):

                col1, col2 = st.columns(2)

                with col1:
                    train_number = st.text_input(
                        "Номер поезда",
                        value=str(row["train_number"]),
                    )

                    current_date = pd.to_datetime(
                        row["departure_date"]
                    ).date()

                    departure_date = st.date_input(
                        "Дата отправления",
                        value=current_date,
                    )

                    departure_station = st.selectbox(
                        "Станция отправления",
                        [
                            "Минск",
                            "Брест",
                            "Витебск",
                            "Гомель",
                            "Москва",
                        ],
                        index=(
                            [
                                "Минск",
                                "Брест",
                                "Витебск",
                                "Гомель",
                                "Москва",
                            ].index(str(row["departure_station"]))
                            if str(row["departure_station"]) in [
                                "Минск",
                                "Брест",
                                "Витебск",
                                "Гомель",
                                "Москва",
                            ]
                            else 0
                        ),
                    )

                with col2:
                    destination_options = [
                        "Москва",
                        "Брест",
                        "Витебск",
                        "Гомель",
                        "Минск",
                    ]

                    destination_station = st.selectbox(
                        "Станция назначения",
                        destination_options,
                        index=(
                            destination_options.index(
                                str(row["destination_station"])
                            )
                            if str(row["destination_station"])
                            in destination_options
                            else 0
                        ),
                    )

                    sold_tickets = st.number_input(
                        "Продано билетов",
                        min_value=0,
                        value=int(row["sold_tickets"]),
                        step=1,
                    )

                    ticket_price = st.number_input(
                        "Цена билета",
                        min_value=0.01,
                        value=float(row["ticket_price"]),
                        step=0.50,
                    )

                carriage_options = [
                    "Плацкарт",
                    "Купейный",
                    "СВ",
                    "Общий",
                ]

                carriage_type = st.selectbox(
                    "Тип вагона",
                    carriage_options,
                    index=(
                        carriage_options.index(
                            str(row["carriage_type"])
                        )
                        if str(row["carriage_type"])
                        in carriage_options
                        else 0
                    ),
                )

                submitted = st.form_submit_button(
                    "Сохранить изменения",
                    use_container_width=True,
                )

                if submitted:

                    if departure_station == destination_station:
                        st.warning(
                            "Станция отправления и станция назначения "
                            "должны различаться."
                        )

                    else:
                        try:
                            update_train(
                                train_id=int(selected_id),
                                train_number=train_number,
                                departure_date=departure_date,
                                departure_station=departure_station,
                                destination_station=destination_station,
                                sold_tickets=sold_tickets,
                                ticket_price=ticket_price,
                                carriage_type=carriage_type,
                            )

                            st.success(
                                "Данные поезда успешно изменены."
                            )
                            refresh_data()

                        except IntegrityError as error:
                            handle_db_error(
                                error,
                                "изменение поезда",
                            )

    with tab_delete:

        if trains.empty:
            st.info("Нет поездов для удаления.")
        else:
            selected_id = st.selectbox(
                "Выберите поезд для удаления",
                trains["id"].astype(int).tolist(),
                format_func=lambda x: format_train_label(
                    trains[trains["id"] == x].iloc[0]
                ),
                key="delete_train_id",
            )

            st.warning(
                "Удаление поезда невозможно, если к нему привязаны "
                "пассажиры."
            )

            if st.button(
                "Удалить поезд",
                type="primary",
                key="delete_train_button",
            ):
                try:
                    delete_train(selected_id)
                    st.success("Поезд удалён.")
                    refresh_data()

                except IntegrityError as error:
                    st.error(
                        "Нельзя удалить поезд, потому что к нему "
                        "привязаны пассажиры."
                    )


# ============================================================
# ПАССАЖИРЫ
# ============================================================

elif page == "Пассажиры":

    st.header("Пассажиры")

    tab_view, tab_add, tab_edit, tab_delete = st.tabs(
        ["Просмотр", "Добавить", "Изменить", "Удалить"]
    )

    with tab_view:

        filtered = passengers.copy()

        # Кнопка сортировки находится в одной строке
        # непосредственно справа от панели фильтрации.
        if "passengers_sort_desc" not in st.session_state:
            st.session_state["passengers_sort_desc"] = False

        filter_col, sort_col = st.columns([20, 1], gap="small")

        with filter_col:
            with st.expander("›  Поиск пассажиров", expanded=False):

                search = st.text_input(
                    "Поиск пассажира",
                    placeholder=(
                        "ФИО, номер поезда, станция "
                        "или направление..."
                    ),
                    key="passenger_search",
                    label_visibility="collapsed",
                )

                train_filter = st.selectbox(
                    "Поезд",
                    ["Все"] + sorted(
                        passengers["train_number"]
                        .dropna()
                        .astype(str)
                        .unique()
                        .tolist()
                    ),
                    key="passenger_train_filter",
                )

        with sort_col:
            if st.button(
                "↕",
                key="passengers_sort_button",
                help="Сортировка по ID: по возрастанию / по убыванию",
                use_container_width=True,
            ):
                st.session_state["passengers_sort_desc"] = not st.session_state[
                    "passengers_sort_desc"
                ]
                st.rerun()

        if search.strip():
            mask = (
                safe_contains(
                    filtered["full_name"],
                    search.strip(),
                )
                | safe_contains(
                    filtered["train_number"],
                    search.strip(),
                )
                | safe_contains(
                    filtered["departure_station"],
                    search.strip(),
                )
                | safe_contains(
                    filtered["destination_station"],
                    search.strip(),
                )
            )

            filtered = filtered[mask]

        if train_filter != "Все":
            filtered = filtered[
                filtered["train_number"].astype(str)
                == train_filter
            ]

        # Фильтрация не меняет базовый порядок ID.
        filtered = filtered.sort_values(
            by="id",
            ascending=not st.session_state["passengers_sort_desc"],
        )

        st.write(f"Найдено пассажиров: **{len(filtered)}**")

        display_passengers = display_date_columns(filtered)

        st.dataframe(
            display_passengers,
            use_container_width=True,
            hide_index=True,
        )

        excel_download(
            filtered,
            "passengers_report.xlsx",
            "Пассажиры",
        )

    with tab_add:

        st.subheader("Добавление пассажира")

        train_options = cached_train_options()

        if train_options.empty:
            st.warning(
                "Сначала добавьте хотя бы один поезд."
            )
        else:
            # Поезда в списке всегда идут по ID: 1, 2, 3, 4...
            train_options = train_options.sort_values(
                by="id",
                ascending=True,
            ).reset_index(drop=True)

            train_ids = train_options["id"].astype(int).tolist()

            # Выбор поезда находится вне формы, поэтому маршрут
            # автоматически обновляется сразу после выбора поезда.
            selected_train_id = st.selectbox(
                "Поезд",
                train_ids,
                format_func=lambda x: format_train_label(
                    train_options[
                        train_options["id"] == x
                    ].iloc[0]
                ),
                key="add_passenger_train",
            )

            selected_train = train_options[
                train_options["id"] == selected_train_id
            ].iloc[0]

            departure_station = str(selected_train["departure_station"])
            destination_station = str(selected_train["destination_station"])

            st.markdown(
                f"**Маршрут:** {departure_station} → {destination_station}"
            )

            st.caption(
                "Маршрут определяется автоматически выбранным поездом."
            )

            with st.form("add_passenger_form"):

                full_name = st.text_input(
                    "ФИО пассажира",
                    placeholder="Иванов Иван Иванович",
                )

                submitted = st.form_submit_button(
                    "Добавить пассажира",
                    use_container_width=True,
                )

                if submitted:

                    if not full_name.strip():
                        st.warning("Введите ФИО пассажира.")

                    else:
                        try:
                            create_passenger(
                                full_name=full_name,
                                train_id=selected_train_id,
                            )

                            st.success(
                                "Пассажир успешно добавлен."
                            )
                            refresh_data()

                        except IntegrityError as error:
                            handle_db_error(
                                error,
                                "добавление пассажира",
                            )

    with tab_edit:

        if passengers.empty:
            st.info("Нет пассажиров для изменения.")
        else:
            passenger_ids = passengers["id"].astype(int).tolist()

            selected_id = st.selectbox(
                "Выберите пассажира",
                passenger_ids,
                format_func=lambda x: (
                    f'{passengers[passengers["id"] == x].iloc[0]["full_name"]} '
                    f'— '
                    f'{passengers[passengers["id"] == x].iloc[0]["train_number"]}'
                ),
                key="edit_passenger_id",
            )

            row = passengers[
                passengers["id"] == selected_id
            ].iloc[0]

            train_options = cached_train_options()
            train_ids = train_options["id"].astype(int).tolist()

            current_train_id = int(row["train_id"])

            if current_train_id not in train_ids:
                st.error(
                    "Текущий поезд пассажира отсутствует в таблице поездов."
                )
            else:

                with st.form("edit_passenger_form"):

                    full_name = st.text_input(
                        "ФИО пассажира",
                        value=str(row["full_name"]),
                    )

                    selected_train_id = st.selectbox(
                        "Поезд",
                        train_ids,
                        index=train_ids.index(current_train_id),
                        format_func=lambda x: format_train_label(
                            train_options[
                                train_options["id"] == x
                            ].iloc[0]
                        ),
                    )

                    selected_train = train_options[
                        train_options["id"] == selected_train_id
                    ].iloc[0]

                    departure_station = selected_train[
                        "departure_station"
                    ]

                    destination_station = selected_train[
                        "destination_station"
                    ]

                    col1, col2 = st.columns(2)

                    with col1:
                        st.text_input(
                            "Станция отправления",
                            value=str(departure_station),
                            disabled=True,
                        )

                    with col2:
                        st.text_input(
                            "Станция назначения",
                            value=str(destination_station),
                            disabled=True,
                        )

                    st.caption(
                        "При смене поезда маршрут автоматически изменится "
                        "на маршрут нового поезда."
                    )

                    submitted = st.form_submit_button(
                        "Сохранить изменения",
                        use_container_width=True,
                    )

                    if submitted:

                        if not full_name.strip():
                            st.warning("Введите ФИО пассажира.")

                        else:
                            try:
                                update_passenger(
                                    passenger_id=selected_id,
                                    full_name=full_name,
                                    train_id=selected_train_id,
                                )

                                st.success(
                                    "Данные пассажира изменены."
                                )
                                refresh_data()

                            except IntegrityError as error:
                                handle_db_error(
                                    error,
                                    "изменение пассажира",
                                )

    with tab_delete:

        if passengers.empty:
            st.info("Нет пассажиров для удаления.")
        else:
            selected_id = st.selectbox(
                "Выберите пассажира для удаления",
                passengers["id"].astype(int).tolist(),
                format_func=lambda x: (
                    f'{passengers[passengers["id"] == x].iloc[0]["full_name"]} '
                    f'— '
                    f'{passengers[passengers["id"] == x].iloc[0]["train_number"]}'
                ),
                key="delete_passenger_id",
            )

            if st.button(
                "Удалить пассажира",
                type="primary",
                key="delete_passenger_button",
            ):
                try:
                    delete_passenger(selected_id)
                    st.success(
                        "Пассажир удалён. Связанный багаж также удалён."
                    )
                    refresh_data()

                except IntegrityError as error:
                    handle_db_error(
                        error,
                        "удаление пассажира",
                    )


# ============================================================
# БАГАЖ
# ============================================================

elif page == "Багаж":

    st.header("Багаж")

    tab_view, tab_add, tab_edit, tab_delete = st.tabs(
        ["Просмотр", "Добавить", "Изменить", "Удалить"]
    )

    with tab_view:

        col1, col2 = st.columns(2)

        with col1:
            st.metric(
                "Всего единиц багажа",
                int(luggage["luggage_count"].sum())
                if not luggage.empty else 0,
            )

        with col2:
            st.metric(
                "Дополнительные места",
                int(luggage["extra_places"].sum())
                if not luggage.empty else 0,
            )

        filtered = luggage.copy()

        search = st.text_input(
            "Поиск багажа",
            placeholder="ФИО, поезд, станция...",
            key="luggage_search",
        )

        if search.strip():
            mask = (
                safe_contains(
                    filtered["full_name"],
                    search.strip(),
                )
                | safe_contains(
                    filtered["train_number"],
                    search.strip(),
                )
                | safe_contains(
                    filtered["departure_station"],
                    search.strip(),
                )
                | safe_contains(
                    filtered["destination_station"],
                    search.strip(),
                )
            )

            filtered = filtered[mask]

        st.write(f"Найдено записей: **{len(filtered)}**")

        st.dataframe(
            display_date_columns(filtered),
            use_container_width=True,
            hide_index=True,
        )

        excel_download(
            filtered,
            "luggage_report.xlsx",
            "Багаж",
        )

    with tab_add:

        st.subheader("Добавление багажа")

        passenger_options = cached_passenger_options()

        if passenger_options.empty:
            st.warning(
                "Сначала добавьте пассажира."
            )
        else:
            passenger_ids = (
                passenger_options["id"]
                .astype(int)
                .tolist()
            )

            with st.form("add_luggage_form"):

                selected_passenger_id = st.selectbox(
                    "Пассажир",
                    passenger_ids,
                    format_func=lambda x: format_passenger_label(
                        passenger_options[
                            passenger_options["id"] == x
                        ].iloc[0]
                    ),
                )

                luggage_count = st.number_input(
                    "Количество багажа",
                    min_value=0,
                    value=0,
                    step=1,
                )

                extra_places = st.number_input(
                    "Дополнительные места",
                    min_value=0,
                    value=0,
                    step=1,
                )

                submitted = st.form_submit_button(
                    "Добавить багаж",
                    use_container_width=True,
                )

                if submitted:
                    try:
                        create_luggage(
                            passenger_id=selected_passenger_id,
                            luggage_count=luggage_count,
                            extra_places=extra_places,
                        )

                        st.success("Багаж добавлен.")
                        refresh_data()

                    except IntegrityError as error:
                        handle_db_error(
                            error,
                            "добавление багажа",
                        )

    with tab_edit:

        if luggage.empty:
            st.info("Нет записей багажа для изменения.")
        else:

            selected_id = st.selectbox(
                "Выберите запись багажа",
                luggage["id"].astype(int).tolist(),
                format_func=lambda x: (
                    f'#{x} — '
                    f'{luggage[luggage["id"] == x].iloc[0]["full_name"]} — '
                    f'{luggage[luggage["id"] == x].iloc[0]["train_number"]}'
                ),
                key="edit_luggage_id",
            )

            row = luggage[
                luggage["id"] == selected_id
            ].iloc[0]

            passenger_options = cached_passenger_options()
            passenger_ids = (
                passenger_options["id"]
                .astype(int)
                .tolist()
            )

            current_passenger_id = int(row["passenger_id"])

            with st.form("edit_luggage_form"):

                selected_passenger_id = st.selectbox(
                    "Пассажир",
                    passenger_ids,
                    index=(
                        passenger_ids.index(
                            current_passenger_id
                        )
                        if current_passenger_id in passenger_ids
                        else 0
                    ),
                    format_func=lambda x: format_passenger_label(
                        passenger_options[
                            passenger_options["id"] == x
                        ].iloc[0]
                    ),
                )

                luggage_count = st.number_input(
                    "Количество багажа",
                    min_value=0,
                    value=int(row["luggage_count"]),
                    step=1,
                )

                extra_places = st.number_input(
                    "Дополнительные места",
                    min_value=0,
                    value=int(row["extra_places"]),
                    step=1,
                )

                submitted = st.form_submit_button(
                    "Сохранить изменения",
                    use_container_width=True,
                )

                if submitted:
                    try:
                        update_luggage(
                            luggage_id=selected_id,
                            passenger_id=selected_passenger_id,
                            luggage_count=luggage_count,
                            extra_places=extra_places,
                        )

                        st.success("Данные багажа изменены.")
                        refresh_data()

                    except IntegrityError as error:
                        handle_db_error(
                            error,
                            "изменение багажа",
                        )

    with tab_delete:

        if luggage.empty:
            st.info("Нет записей багажа для удаления.")
        else:

            selected_id = st.selectbox(
                "Выберите запись багажа",
                luggage["id"].astype(int).tolist(),
                format_func=lambda x: (
                    f'#{x} — '
                    f'{luggage[luggage["id"] == x].iloc[0]["full_name"]}'
                ),
                key="delete_luggage_id",
            )

            if st.button(
                "Удалить багаж",
                type="primary",
                key="delete_luggage_button",
            ):
                try:
                    delete_luggage(selected_id)
                    st.success("Запись багажа удалена.")
                    refresh_data()

                except IntegrityError as error:
                    handle_db_error(
                        error,
                        "удаление багажа",
                    )


# ============================================================
# МАРШРУТЫ
# ============================================================

elif page == "Маршруты":

    st.header("Аналитика маршрутов")

    routes = cached_routes()

    if routes.empty:
        st.info("Нет данных для анализа маршрутов.")
    else:
        st.dataframe(
            routes,
            use_container_width=True,
            hide_index=True,
        )

        excel_download(
            routes,
            "routes_report.xlsx",
            "Маршруты",
            "Скачать Excel-отчёт",
        )


# ============================================================
# АНАЛИТИКА
# ============================================================

elif page == "Аналитика":

    st.header("Расширенная аналитика")

    tab1, tab2, tab3 = st.tabs(
        ["По датам", "Типы вагонов", "KPI"]
    )

    with tab1:

        daily = cached_daily()

        if daily.empty:
            st.info("Нет данных.")
        else:
            display_daily = display_date_columns(daily)

            st.dataframe(
                display_daily,
                use_container_width=True,
                hide_index=True,
            )

            chart_daily = daily.copy()
            chart_daily["departure_date"] = pd.to_datetime(
                chart_daily["departure_date"],
                errors="coerce",
            )

            st.line_chart(
                chart_daily.set_index("departure_date")[
                    ["total_tickets", "total_revenue"]
                ]
            )

            excel_download(
                display_daily,
                "daily_analytics_report.xlsx",
                "По датам",
                "Скачать Excel-отчёт",
            )

    with tab2:

        carriages = cached_carriages()

        if carriages.empty:
            st.info("Нет данных.")
        else:
            st.dataframe(
                carriages,
                use_container_width=True,
                hide_index=True,
            )

            st.bar_chart(
                carriages.set_index("carriage_type")[
                    "total_revenue"
                ]
            )

            excel_download(
                carriages,
                "carriage_analytics_report.xlsx",
                "Типы вагонов",
                "Скачать Excel-отчёт",
            )

    with tab3:

        col1, col2, col3 = st.columns(3)

        with col1:
            st.metric(
                "Средняя цена билета",
                f'{summary["average_ticket_price"]:.2f} BYN',
            )

        with col2:
            st.metric(
                "Всего багажа",
                summary["luggage"],
            )

        with col3:
            average_revenue = (
                summary["revenue"] / summary["trains"]
                if summary["trains"] > 0
                else 0
            )

            st.metric(
                "Средняя выручка на поезд",
                f"{average_revenue:,.2f} BYN",
            )
